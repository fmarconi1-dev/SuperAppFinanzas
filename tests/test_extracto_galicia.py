"""Lector del extracto de Galicia (Fase D, paso 4).

Los Excel se arman en el momento con `openpyxl`, imitando el formato real:
filas de encabezado, títulos con acentos rotos, todo como texto, lo más nuevo
primero. **Los nombres son inventados.** El test contra el archivo real de
Franco se saltea si `data/extractos/` no está, y solo mira cantidades.
"""

from __future__ import annotations

import io
from collections import Counter
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from places.finanzas.core.extracto_galicia import (
    ExtractoInvalido,
    centavos,
    leer_extracto,
    normalizar_tipo,
)

ENCABEZADO = [
    ["Banco Galicia - Caja Ahorro"],
    ["Nro. de Cuenta: ...000000"],
    ["Fecha Actual: 18/11/2026"],
    ["Hora Actual: 10:00"],
    ["Intervalo de Consulta: de 17/10/2026 a 16/11/2026"],
    ["Fecha", "Movimiento", "D�bito", "Cr�dito", "Saldo Parcial", "Comentarios"],
]


def armar_extracto(movimientos, saldo_inicial="100000,00") -> bytes:
    """`movimientos` de la más vieja a la más nueva: (fecha, texto, monto en pesos).

    Calcula los saldos y los escribe como el banco: lo más nuevo arriba, los
    montos con puntos de miles y el saldo sin ellos.
    """
    saldo = centavos(saldo_inicial, que="saldo", fila=0)
    filas = []
    for fecha, texto, monto in movimientos:
        c = round(monto * 100)
        saldo += c
        filas.append([fecha, texto, _pesos(c) if c < 0 else "0,00",
                      _pesos(c) if c > 0 else "0,00", _saldo(saldo), ""])
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Cuentas"
    for f in ENCABEZADO + list(reversed(filas)):
        hoja.append(f)
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _pesos(c: int) -> str:
    entero, dec = divmod(abs(c), 100)
    return f"{'-' if c < 0 else ''}{entero:,}".replace(",", ".") + f",{dec:02d}"


def _saldo(c: int) -> str:
    entero, dec = divmod(abs(c), 100)
    return f"{'-' if c < 0 else ''}{entero},{dec:02d}"


COMPRA = "COMPRA DEBITO\nSUPERMERCADO INVENTADO\n4517XXXXXXXXXX12\nA123"
TRANSFER = "TRANSFERENCIA A TERCEROS\n20123456789\n0000003100012345678901\nJUANA INVENTADA\nVARIOS\nBANCO X"


def test_lee_y_da_vuelta_el_orden():
    datos = armar_extracto([
        ("17/10/2026", COMPRA, -2500),
        ("17/10/2026", "SUSCRIPCION FIMA\nFIMA PREMIUM CLASE A\nNro Operacion: 123", -50000),
        ("18/10/2026", "TRANSFERENCIA DE TERCEROS\nPEDRO INVENTADO\n20111111112\nVARIOS", 25000),
    ])
    e = leer_extracto(datos)
    assert [f.tipo for f in e.filas] == ["COMPRA DEBITO", "SUSCRIPCION FIMA",
                                          "TRANSFERENCIA DE TERCEROS"]
    assert [f.monto for f in e.filas] == [-250000, -5000000, 2500000]
    assert (e.desde, e.hasta) == (date(2026, 10, 17), date(2026, 10, 18))
    assert e.filas[0].saldo_antes == 10000000
    assert e.saldo_final == 10000000 - 250000 - 5000000 + 2500000


def test_el_comercio_es_la_primera_linea_de_texto():
    e = leer_extracto(armar_extracto([("17/10/2026", COMPRA, -100),
                                      ("17/10/2026", TRANSFER, -200)]))
    assert e.filas[0].comercio == "SUPERMERCADO INVENTADO"
    assert e.filas[1].comercio == "JUANA INVENTADA"     # después del CUIT y el CBU


def test_los_tipos_sin_comercio_no_lo_inventan():
    e = leer_extracto(armar_extracto([
        ("17/10/2026", "SUSCRIPCION FIMA\nFIMA PREMIUM CLASE A", -100),
        ("17/10/2026", "PAGO TARJETA VISA\nOPERACION 1234567890", -200),
        ("17/10/2026", "IVA", -10),
    ]))
    assert [f.comercio for f in e.filas] == [None, None, None]


def test_nunca_lee_cuit_ni_cbu():
    e = leer_extracto(armar_extracto([("17/10/2026", TRANSFER, -200)]))
    texto = repr(e)
    assert "20123456789" not in texto and "0000003100012345678901" not in texto


def test_la_cotizacion_de_los_dolares():
    e = leer_extracto(armar_extracto([
        ("13/10/2026", "Compra venta de dolares\nFRANCO INVENTADO\n20999999990\n"
                       "Cotizacion: 1535.00000000000000\nOPERACION 1\n"
                       "Compra Moneda Extranjera: D�lares", -153500)]))
    assert e.filas[0].tipo == "COMPRA VENTA DE DOLARES"
    assert e.filas[0].cotizacion == 1535.0


@pytest.mark.parametrize("crudo, tipo", [
    ("DEP.EFVO.AUTOSERVICIO TICKET:  109477", "DEP.EFVO.AUTOSERVICIO TICKET:"),
    ("Compra venta de dolares", "COMPRA VENTA DE DOLARES"),
    ("PERCEPCION RG 5617/24", "PERCEPCION RG 5617/24"),
    ("  TRANSF. CTAS   PROPIAS ", "TRANSF. CTAS PROPIAS"),
    ("Comisión", "COMISION"),
])
def test_normalizar_tipo(crudo, tipo):
    assert normalizar_tipo(crudo) == tipo


@pytest.mark.parametrize("texto, valor", [
    ("-200.000,00", -20000000),
    ("153263,01", 15326301),
    ("709269,8", 70926980),       # el saldo a veces viene con un decimal
    ("0,00", 0),
    ("25.000,00", 2500000),
    ("1.250.000,50", 125000050),
])
def test_centavos(texto, valor):
    assert centavos(texto, que="monto", fila=1) == valor


@pytest.mark.parametrize("texto", ["2500.50", "abc", "", "1.5,00", "-"])
def test_centavos_estricto(texto):
    with pytest.raises(ExtractoInvalido):
        centavos(texto, que="monto", fila=1)


def test_una_cadena_rota_se_rechaza():
    datos = armar_extracto([("17/10/2026", COMPRA, -100), ("18/10/2026", COMPRA, -200)])
    libro = openpyxl.load_workbook(io.BytesIO(datos))
    libro.active["E7"] = "1,00"          # el saldo de la fila más nueva, adulterado
    salida = io.BytesIO()
    libro.save(salida)
    with pytest.raises(ExtractoInvalido, match="cadena de saldos se corta el 18/10/2026"):
        leer_extracto(salida.getvalue())


def test_un_archivo_que_no_es_excel():
    with pytest.raises(ExtractoInvalido, match="No pude abrir el Excel"):
        leer_extracto(b"%PDF-1.7 no soy un excel")


def test_un_excel_sin_movimientos():
    libro = openpyxl.Workbook()
    libro.active.append(["hola"])
    salida = io.BytesIO()
    libro.save(salida)
    with pytest.raises(ExtractoInvalido, match="No encontré movimientos"):
        leer_extracto(salida.getvalue())


# ---------------------------------------------------------------- el archivo real

_REAL = Path(__file__).resolve().parents[1] / "data" / "extractos"
_ARCHIVOS = sorted(_REAL.glob("galicia-caja-ahorro-*.xlsx")) if _REAL.is_dir() else []


@pytest.mark.skipif(not _ARCHIVOS, reason="sin el extracto de Franco (data/extractos)")
def test_el_extracto_real_de_franco():
    """406 filas, la cadena no se corta, y al 16/10/2026 la caja tenía $153.263,01."""
    e = leer_extracto(_ARCHIVOS[0].read_bytes())
    assert len(e.filas) == 406
    assert e.hasta == date(2026, 10, 16) and e.saldo_final == 15326301
    tipos = Counter(f.tipo for f in e.filas)
    assert tipos["COMPRA DEBITO"] > 100
    assert all(f.comercio for f in e.filas if f.tipo == "COMPRA DEBITO")
