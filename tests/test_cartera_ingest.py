"""Tests del ingest de cartera (Fase 3a): carga de instrumentos + movimientos.

Sintético y autocontenido (corre en CI): no depende del Excel real. La
reconciliación contra la planilla vive en `test_cartera_recon.py`.
"""

from __future__ import annotations

import pytest

from core import db as core_db
from core.db import connect, init_db
from places.inversiones.cartera.ingest import (
    cargar_instrumentos_en_db,
    importar_movimientos,
    instrumentos_desde_config,
)
from places.inversiones.cartera.parser import Instrumento, parse_filas
from places.inversiones.cartera.posiciones import posiciones_desde_db


@pytest.fixture()
def db(tmp_path, monkeypatch):
    db_path = tmp_path / "finanzas.db"
    monkeypatch.setenv("DB_PATH", str(db_path))
    monkeypatch.setattr(core_db, "BACKUP_DIR", tmp_path / "backups")
    init_db(db_path)
    conn = connect(db_path)
    conn.execute(
        "INSERT OR IGNORE INTO usuarios (id, username, fullname) "
        "VALUES (1, 'franco', 'Franco')"
    )
    yield conn
    conn.close()


@pytest.fixture()
def instrumentos():
    I = Instrumento
    return {
        "GGAL": I("GGAL", "GGALD", "Acciones", "ARS"),
        "SPY": I("SPY", "SPYD", "Cedears", "ARS"),
    }


def _movs(instrumentos):
    filas = [
        ["Compra", "2026-01-10", "2026-01-12", "GGAL | Grupo", -1000, 10, 100, "ARS"],
        ["Compra", "2026-02-10", "2026-02-12", "GGAL | Grupo", -1100, 5, 110, "ARS"],
        ["Venta", "2026-03-10", "2026-03-12", "GGAL | Grupo", 700, -4, 115, "ARS"],
        ["Compra", "2026-01-21", "2026-01-23", "SPYD | SPDR", -176.79, 5, 35.23, "USD"],
    ]
    return parse_filas(filas, instrumentos)


def _mini_excel(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Config"
    headers = ["Ticker ARS", "Ticker USD", "Descripción", "Categoría",
               "Panel Precio", "Activo en cartera", "Moneda Cotiz", "Divisor a USD"]
    for i, h in enumerate(headers):
        ws.cell(4, 7 + i, h)
    filas = [
        ("GGAL", "GGALD", "Grupo Financiero Galicia", "Acciones", "PANEL LIDER", "SI", "ARS", 1408),
        ("AE38", "AE38D", "Bono USD 2038", "Bonos", "BONOS", "SI", "USD/100", 100),
    ]
    for r, fila in enumerate(filas, start=5):
        for i, v in enumerate(fila):
            ws.cell(r, 7 + i, v)
    p = tmp_path / "mini.xlsx"
    wb.save(p)
    return p


# ---------------------------------------------------------------------------

def test_importar_y_dedupe(db, instrumentos):
    ins, skip = importar_movimientos(db, _movs(instrumentos), user_id=1)
    assert (ins, skip) == (4, 0)
    # re-importar el mismo archivo → todo saltado por el hash.
    ins2, skip2 = importar_movimientos(db, _movs(instrumentos), user_id=1)
    assert (ins2, skip2) == (0, 4)
    n = db.execute("SELECT COUNT(*) c FROM cartera_movimientos").fetchone()["c"]
    assert n == 4


def test_posiciones_desde_db(db, instrumentos):
    importar_movimientos(db, _movs(instrumentos), user_id=1)
    pos = posiciones_desde_db(db, user_id=1)
    assert pos["GGAL"].cant_actual == 11        # 10 + 5 - 4
    assert pos["SPY"].cant_actual == 5          # SPYD pliega a SPY (Config)
    assert pos["GGAL"].costo_ars_compras == pytest.approx(2100.0)  # 1000 + 1100


def test_ingest_aisla_por_usuario(db, instrumentos):
    importar_movimientos(db, _movs(instrumentos), user_id=1)
    # mismo archivo, otro usuario: entra (hash único por user_id).
    ins, _ = importar_movimientos(db, _movs(instrumentos), user_id=2)
    assert ins == 4
    assert "GGAL" in posiciones_desde_db(db, user_id=2)


def test_instrumentos_desde_config(tmp_path):
    insts = instrumentos_desde_config(_mini_excel(tmp_path))
    assert insts["GGAL"].ticker_usd == "GGALD"
    assert insts["AE38"].moneda_cotiz == "USD/100"
    assert insts["AE38"].divide_por_100 is True


def test_cargar_instrumentos_en_db(db, tmp_path):
    n = cargar_instrumentos_en_db(db, _mini_excel(tmp_path), user_id=1)
    db.commit()
    assert n == 2
    row = db.execute(
        "SELECT * FROM cartera_instrumentos WHERE ticker='AE38' AND user_id=1"
    ).fetchone()
    assert row["moneda_cotiz"] == "USD/100"
    assert row["divisor_usd"] == 100
    assert row["categoria_activo"] == "Bonos"


# ---------------------------------------------------------------------------
# Carga en un paso: detección del tipo de archivo
# ---------------------------------------------------------------------------

def _excel_propio(tmp_path):
    """Un archivo con las hojas Config y Movimientos Crudo (el de Franco)."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Config"
    for i, h in enumerate(["Ticker ARS", "Ticker USD", "Descripción", "Categoría",
                           "Panel", "Activo", "Moneda Cotiz", "Divisor"]):
        ws.cell(4, 7 + i, h)
    for i, v in enumerate(("GGAL", "GGALD", "Galicia", "Acciones", "PANEL LIDER",
                           "SI", "ARS", 1408)):
        ws.cell(5, 7 + i, v)
    ws2 = wb.create_sheet("Movimientos Crudo")
    for i, h in enumerate(["Operación", "Concertación", "Liquidación",
                           "Descripción", "Monto", "Cantidad", "Precio", "Moneda"]):
        ws2.cell(1, 1 + i, h)
    p = tmp_path / "Inversiones.xlsx"
    wb.save(p)
    return p


def _export_movimientos(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "-movimientos"
    for i, h in enumerate(["Operación", "Concertación", "Liquidación",
                           "Descripción", "Monto", "Cantidad", "Precio", "Moneda"]):
        ws.cell(1, 1 + i, h)
    p = tmp_path / "inviu-movimientos.xlsx"
    wb.save(p)
    return p


def _export_tenencia(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(2, 1, "Fecha de valuación:")
    ws.cell(2, 2, "30/07/2026")
    ws.cell(3, 1, "TC USD MEP:")
    ws.cell(3, 2, 1522.185)
    p = tmp_path / "Tenencias.xlsx"
    wb.save(p)
    return p


def test_detecta_los_tres_tipos_por_contenido(tmp_path):
    """Franco sube los tres juntos sin decir cuál es cuál: la app los reconoce
    por su estructura, no por el nombre del archivo."""
    from places.inversiones.cartera.ingest import (
        TIPO_EXCEL_PROPIO, TIPO_MOVIMIENTOS, TIPO_TENENCIA, detectar_tipo,
    )
    assert detectar_tipo(_excel_propio(tmp_path)) == TIPO_EXCEL_PROPIO
    assert detectar_tipo(_export_movimientos(tmp_path)) == TIPO_MOVIMIENTOS
    assert detectar_tipo(_export_tenencia(tmp_path)) == TIPO_TENENCIA


def test_archivo_irreconocible_no_rompe(tmp_path):
    import openpyxl
    from places.inversiones.cartera.ingest import TIPO_DESCONOCIDO, detectar_tipo
    wb = openpyxl.Workbook()
    wb.active.cell(1, 1, "cualquier cosa")
    p = tmp_path / "raro.xlsx"
    wb.save(p)
    assert detectar_tipo(p) == TIPO_DESCONOCIDO
    # un archivo que ni se puede abrir tampoco debe explotar
    malo = tmp_path / "roto.xlsx"
    malo.write_bytes(b"no soy un excel")
    assert detectar_tipo(malo) == TIPO_DESCONOCIDO


# ---------------------------------------------------------------------------
# Regresión: importar sin el mapa de instrumentos dejaba tickers sin plegar
# ---------------------------------------------------------------------------

def test_remapear_tickers_base_repara_variantes_usd(db, instrumentos):
    """Bug real en producción (30/7/2026): si los movimientos se importan SIN
    el mapa de instrumentos, las variantes en dólares quedan como posiciones
    propias (SPYD aparte de SPY). Como el dedupe salta las filas ya
    existentes, reimportar con el Excel NO lo corregía: el error quedaba
    grabado, inflando posiciones y costo. En la cartera de Franco: 37
    posiciones en vez de 22."""
    from places.inversiones.cartera.ingest import remapear_tickers_base

    # importado SIN instrumentos → ticker_base = ticker_crudo
    filas = [
        ["Compra", "2026-01-10", "2026-01-12", "SPY | SPDR", -1000, 10, 100, "ARS"],
        ["Compra", "2026-01-21", "2026-01-23", "SPYD | SPDR", -176, 5, 35, "USD"],
    ]
    movs = parse_filas(filas, {})            # sin mapping
    importar_movimientos(db, movs, user_id=1)
    pos = posiciones_desde_db(db, 1)
    assert {"SPY", "SPYD"} <= set(pos)       # el bug: dos posiciones

    # ahora se carga el mapa y se repara
    n = remapear_tickers_base(db, instrumentos, user_id=1)
    assert n == 1                            # solo la fila de SPYD
    pos2 = posiciones_desde_db(db, 1, instrumentos)
    assert "SPYD" not in pos2                # plegada
    assert pos2["SPY"].cant_comprada == 15   # 10 + 5


def test_remapear_es_idempotente(db, instrumentos):
    from places.inversiones.cartera.ingest import remapear_tickers_base
    movs = parse_filas(
        [["Compra", "2026-01-21", "2026-01-23", "SPYD | SPDR", -176, 5, 35, "USD"]],
        {})
    importar_movimientos(db, movs, user_id=1)
    assert remapear_tickers_base(db, instrumentos, user_id=1) == 1
    assert remapear_tickers_base(db, instrumentos, user_id=1) == 0   # ya está


def test_remapear_sin_instrumentos_no_hace_nada(db):
    from places.inversiones.cartera.ingest import remapear_tickers_base
    assert remapear_tickers_base(db, {}, user_id=1) == 0


def test_instrumentos_de_db_lee_el_flag_activo(db, tmp_path):
    """Segundo bug del 30/7: `_instrumentos_de_db` no leía la columna `activo`,
    así que TODO quedaba como vigente y las letras/bonos vencidos seguían
    figurando con su cantidad residual (el vencimiento entra como
    AMORTIZACION y no descuenta unidades)."""
    from places.inversiones.cartera.view import _instrumentos_de_db

    db.execute(
        "INSERT INTO cartera_instrumentos (ticker, user_id, categoria_activo, "
        " moneda_cotiz, activo) VALUES ('S15G5', 1, 'Letras', 'ARS', 0)")
    db.execute(
        "INSERT INTO cartera_instrumentos (ticker, user_id, categoria_activo, "
        " moneda_cotiz, activo) VALUES ('GGAL', 1, 'Acciones', 'ARS', 1)")
    db.commit()
    insts = _instrumentos_de_db(db, 1)
    assert insts["S15G5"].activo is False     # vencida
    assert insts["GGAL"].activo is True


# ---------------------------------------------------------------------------
# Corrección de un movimiento mal informado por el broker
# ---------------------------------------------------------------------------

def _filas_asml(monto, precio):
    """La compra real de Franco, con el error del broker (÷100) o corregida."""
    return [
        ["Operación", "Concertación", "Liquidación", "Descripción", "Monto",
         "Cantidad", "Precio", "Moneda"],
        ["Compra", "2025-05-16", "2025-05-19", "ASML | ASML HOLDING NV",
         monto, 9, precio, "ARS"],
    ]


def test_corregir_movimiento_deja_el_hash_que_produciria_el_excel_corregido(
    db, instrumentos
):
    """El corazón del asunto: después de corregir en la app y corregir el
    Excel con los mismos valores, **reimportar no debe duplicar**."""
    from places.inversiones.cartera.ingest import (
        corregir_movimiento, movimientos_de_ticker,
    )

    insts = dict(instrumentos)
    insts["ASML"] = Instrumento(ticker="ASML", ticker_usd="ASMLD",
                                categoria_activo="Cedears")
    importar_movimientos(db, parse_filas(_filas_asml(-543.05, 60.1), insts), 1)
    db.commit()

    (mov,) = movimientos_de_ticker(db, 1, "ASML")
    assert corregir_movimiento(db, 1, mov["id"], monto=-54305.0, precio=6010.0)
    db.commit()

    # Reimportar el Excel YA corregido: el dedupe tiene que reconocerlo.
    insertados, saltados = importar_movimientos(
        db, parse_filas(_filas_asml(-54305.0, 6010.0), insts), 1)
    assert (insertados, saltados) == (0, 1), "se duplicó al reimportar"
    filas = movimientos_de_ticker(db, 1, "ASML")
    assert len(filas) == 1
    assert filas[0]["monto"] == pytest.approx(-54305.0)
    assert filas[0]["cantidad"] == pytest.approx(9)


def test_sin_corregir_en_la_app_el_excel_arreglado_DUPLICA(db, instrumentos):
    """La trampa, congelada en un test: arreglar solo el Excel no alcanza.

    Si algún día alguien decide que la corrección desde la app "sobra", este
    test muestra qué pasa — quedan las dos compras, 18 unidades de ASML donde
    hay 9.
    """
    from places.inversiones.cartera.ingest import movimientos_de_ticker

    insts = dict(instrumentos)
    insts["ASML"] = Instrumento(ticker="ASML", categoria_activo="Cedears")
    importar_movimientos(db, parse_filas(_filas_asml(-543.05, 60.1), insts), 1)
    db.commit()

    insertados, _ = importar_movimientos(
        db, parse_filas(_filas_asml(-54305.0, 6010.0), insts), 1)
    assert insertados == 1
    filas = movimientos_de_ticker(db, 1, "ASML")
    assert len(filas) == 2
    assert sum(f["cantidad"] for f in filas) == 18


def test_corregir_movimiento_sin_cambios_no_toca_nada(db, instrumentos):
    from places.inversiones.cartera.ingest import (
        corregir_movimiento, movimientos_de_ticker,
    )

    insts = dict(instrumentos)
    insts["ASML"] = Instrumento(ticker="ASML", categoria_activo="Cedears")
    importar_movimientos(db, parse_filas(_filas_asml(-543.05, 60.1), insts), 1)
    db.commit()
    (mov,) = movimientos_de_ticker(db, 1, "ASML")
    assert corregir_movimiento(db, 1, mov["id"], monto=-543.05) is False


def test_corregir_movimiento_de_otro_usuario_no_hace_nada(db, instrumentos):
    from places.inversiones.cartera.ingest import (
        corregir_movimiento, movimientos_de_ticker,
    )

    insts = dict(instrumentos)
    insts["ASML"] = Instrumento(ticker="ASML", categoria_activo="Cedears")
    importar_movimientos(db, parse_filas(_filas_asml(-543.05, 60.1), insts), 1)
    db.commit()
    (mov,) = movimientos_de_ticker(db, 1, "ASML")
    assert corregir_movimiento(db, 99, mov["id"], monto=-1.0) is False


def test_corregir_reescala_el_precio_normalizado_de_los_bonos(db, instrumentos):
    """Los bonos cotizan por 100: la corrección no puede perder esa regla."""
    from places.inversiones.cartera.ingest import (
        corregir_movimiento, movimientos_de_ticker,
    )

    insts = {"AE38": Instrumento(ticker="AE38", categoria_activo="Bonos",
                                 moneda_cotiz="USD/100")}
    filas = [
        ["Operación", "Concertación", "Liquidación", "Descripción", "Monto",
         "Cantidad", "Precio", "Moneda"],
        ["Compra", "2025-05-16", "2025-05-19", "AE38 | Bono USD 2038",
         -1000.0, 100, 1000.0, "USD"],
    ]
    importar_movimientos(db, parse_filas(filas, insts), 1)
    db.commit()
    (mov,) = movimientos_de_ticker(db, 1, "AE38")
    antes = db.execute(
        "SELECT precio_normalizado FROM cartera_movimientos WHERE id = ?",
        (mov["id"],)).fetchone()["precio_normalizado"]
    assert antes == pytest.approx(10.0)

    corregir_movimiento(db, 1, mov["id"], monto=-2000.0, precio=2000.0)
    db.commit()
    despues = db.execute(
        "SELECT precio_normalizado FROM cartera_movimientos WHERE id = ?",
        (mov["id"],)).fetchone()["precio_normalizado"]
    assert despues == pytest.approx(20.0)
